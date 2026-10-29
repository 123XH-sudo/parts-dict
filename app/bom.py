from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

from openpyxl import load_workbook

from app.models import Part
from app.search import normalize, parse_aliases


class BomError(ValueError):
    pass


MAX_BYTES = 2 * 1024 * 1024
REQUIRED_FIELDS = ("designator", "name", "footprint")
HEADER_SCAN_ROWS = 10

# 立创EDA 导出的是 Name，嘉立创下单/贴片用的 BOM 是 Comment。
HEADER_ALIASES = {
    "designator": (
        "designator",
        "designators",
        "reference",
        "references",
        "refdes",
        "part reference",
        "位号",
        "标号",
        "元件标号",
        "器件标号",
        "器件位号",
        "元件位号",
    ),
    "name": (
        "name",
        "comment",
        "value",
        "device",
        "规格",
        "规格型号",
        "型号",
        "名称",
        "器件名称",
        "值",
    ),
    "footprint": ("footprint", "package", "封装", "封装形式", "器件封装"),
    "supplier": (
        "supplier",
        "supplier part",
        "supplier part #",
        "supplier part number",
        "lcsc",
        "lcsc part",
        "lcsc part #",
        "lcsc part number",
        "lcsc part no",
        "供应商编号",
        "供应商料号",
        "立创编号",
        "嘉立创编号",
        "商品编号",
    ),
    "quantity": ("quantity", "qty", "数量"),
}


def _header_key(cell: object) -> str:
    text = str(cell).strip().lstrip("\ufeff").lower() if cell is not None else ""
    return re.sub(r"[\s_:：]+", "", text).replace("＃", "#")


_ALIAS_KEYS = {
    field: frozenset(_header_key(alias) for alias in aliases)
    for field, aliases in HEADER_ALIASES.items()
}

SKIP_POLAR_TOKENS = ("test-point", "testpoint", "铜条", "wafer", "conn", "db2erc")
SKIP_PREFIXES = {"R", "C", "L", "TP", "CN", "SW", "X", "NTC"}
POLAR_PREFIXES = {"D", "LED", "Q", "EC"}
POLAR_NAME_TOKENS = ("1n", "ss34", "ss54", "smaj", "bzt", "led", "mos")
POLAR_FP_TOKENS = (
    "sod-",
    "sma_",
    "led",
    "to-220",
    "lqfp",
    "soic",
    "tssop",
    "sot-23",
    "powervdfn",
    "cap-th",
)


@dataclass
class ParsedLine:
    designators: str
    name: str
    footprint: str
    supplier: str
    quantity: int
    skip_bin: bool
    polarized_hint: bool
    warning: str


def name_variants(name: str) -> set[str]:
    base = normalize(name)
    variants = {base}
    if base.endswith("r") and base[:-1].replace(".", "").isdigit():
        variants.add(base[:-1])
    if re.fullmatch(r"[0-9.]+", base):
        variants.add(base + "r")
    return {item for item in variants if item}


def footprint_short(fp: str) -> str:
    compact = fp.strip().replace(" ", "")
    upper = compact.upper()
    matched = re.match(r"^(?:R|C|LED_)(\d{4})$", upper)
    if matched:
        return matched.group(1).lower()
    return normalize(compact)[:20]


def lcsc_code(supplier: str) -> str:
    if not supplier:
        return ""
    first = supplier.split(",")[0].strip()
    if re.fullmatch(r"C\d+", first, re.I):
        return first.lower()
    return ""


def designator_prefixes(designators: str) -> list[str]:
    prefixes = []
    for item in re.split(r"[,，;；\s]+", designators.strip()):
        text = item.strip().upper()
        matched = re.match(r"^([A-Z]+)", text)
        if matched:
            prefixes.append(matched.group(1))
    return prefixes


def is_skip_bin(name: str, footprint: str) -> bool:
    blob = f"{name} {footprint}".lower()
    if "test-point" in blob or "testpoint" in blob.replace("-", ""):
        return True
    if "铜条" in name or "铜条" in footprint:
        return True
    return False


def _blob(name: str, footprint: str) -> str:
    return f"{name} {footprint}".lower()


def polarized_hint(designators: str, name: str, footprint: str) -> bool:
    blob = _blob(name, footprint)
    if any(token in blob for token in SKIP_POLAR_TOKENS):
        return False
    prefixes = designator_prefixes(designators)
    if prefixes and all(prefix in SKIP_PREFIXES for prefix in prefixes):
        return False
    if any(prefix in POLAR_PREFIXES for prefix in prefixes):
        return True
    if any(token in blob for token in POLAR_NAME_TOKENS):
        return True
    if any(token in blob for token in POLAR_FP_TOKENS):
        return True
    if "U" in prefixes:
        return True
    return False


def polarity_warning(designators: str, name: str, footprint: str) -> str:
    prefixes = set(designator_prefixes(designators))
    blob = _blob(name, footprint)
    if "LED" in prefixes or "led" in blob:
        return "别贴反：看封装缺口或丝印极性标。"
    if "Q" in prefixes or "mos" in blob:
        return "别贴反：按封装方向，G/D/S 不要转 180°。"
    if "EC" in prefixes or "cap-th" in blob:
        return "别贴反：负极（壳体条纹/短脚）对丝印阴影；长脚一般是正极。"
    if "D" in prefixes or "1n" in blob or "sod-" in blob or "smaj" in blob or "bzt" in blob:
        return "别贴反：阴极有杠/斜角，对丝印。"
    return "别贴反：1 脚对丝印圆点或缺口。"


def _header_map(header: tuple | list) -> dict[str, list[int]] | None:
    found: dict[str, list[int]] = {}
    for index, cell in enumerate(header):
        key = _header_key(cell)
        if not key:
            continue
        for field, keys in _ALIAS_KEYS.items():
            if key in keys:
                found.setdefault(field, []).append(index)
                break
    if all(field in found for field in REQUIRED_FIELDS):
        return found
    return None


def _cell(row: tuple, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    return str(value).strip()


def _supplier_cell(row: tuple, indices: list[int]) -> str:
    values = [value for value in (_cell(row, index) for index in indices) if value]
    for value in values:
        if lcsc_code(value):
            return value
    return values[0] if values else ""


def _designator_count(designators: str) -> int:
    items = [item for item in re.split(r"[,，;；\s]+", designators.strip()) if item]
    return len(items) or 1


def _header_error(rows: list[tuple]) -> BomError:
    header = next((row for row in rows if row and any(str(cell).strip() for cell in row if cell is not None)), ())
    shown = "、".join(str(cell).strip() for cell in header if cell is not None and str(cell).strip())
    message = "没认出 BOM 表头：需要位号（Designator）、规格（Name 或 Comment）、封装（Footprint）三列。"
    if shown:
        message += f"这个文件的表头是：{shown}。"
    return BomError(message)


def _as_lines(mapping: dict[str, list[int]], body: list[tuple]) -> list[ParsedLine]:
    designator_i = mapping["designator"][0]
    name_i = mapping["name"][0]
    footprint_i = mapping["footprint"][0]
    supplier_is = mapping.get("supplier", [])
    qty_i = mapping.get("quantity", [None])[0]
    lines: list[ParsedLine] = []
    for row in body:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        designators = _cell(row, designator_i)
        name = _cell(row, name_i)
        footprint = _cell(row, footprint_i)
        if not designators or not name:
            continue
        quantity = _designator_count(designators)
        if qty_i is not None:
            qty_raw = _cell(row, qty_i)
            try:
                quantity = int(float(qty_raw)) if qty_raw else quantity
            except ValueError:
                pass
        skip = is_skip_bin(name, footprint)
        polar = False if skip else polarized_hint(designators, name, footprint)
        lines.append(
            ParsedLine(
                designators=designators,
                name=name,
                footprint=footprint,
                supplier=_supplier_cell(row, supplier_is),
                quantity=quantity,
                skip_bin=skip,
                polarized_hint=polar,
                warning=polarity_warning(designators, name, footprint) if polar else "",
            )
        )
    if not lines:
        raise BomError("这个 BOM 里没有可用的料，请确认表头和内容后再导出一次。")
    return lines


def _split_rows(rows: list[tuple]) -> list[ParsedLine]:
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping = _header_map(row)
        if mapping is not None:
            return _as_lines(mapping, rows[index + 1 :])
    raise _header_error(rows)


def parse_bom(content: bytes, filename: str) -> list[ParsedLine]:
    if len(content) > MAX_BYTES:
        raise BomError("文件不能超过 2 MB")
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        rows = [tuple(item.strip() for item in row) for row in csv.reader(io.StringIO(text))]
        if not rows:
            raise BomError("这个文件是空的，请用嘉立创导出的 BOM。")
        return _split_rows(rows)
    if not lower.endswith(".xlsx"):
        raise BomError("只支持嘉立创导出的 xlsx 或 csv BOM，请重新导出后再上传。")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise BomError("这个 xlsx 打不开，请重新从嘉立创导出 BOM。") from exc
    first_error: BomError | None = None
    for name in workbook.sheetnames:
        rows = list(workbook[name].iter_rows(values_only=True))
        if not rows:
            continue
        try:
            return _split_rows(rows)
        except BomError as exc:
            first_error = first_error or exc
    raise first_error or BomError("这个文件里没有表，请用嘉立创导出的 BOM。")


_PKG_CODES = ("0201", "0402", "0603", "0805", "1206", "1210", "1812", "2010", "2512", "0602", "0508")
_PKG_RE = re.compile(r"(?<!\d)(" + "|".join(_PKG_CODES) + r")(?!\d)")


def _part_keys(part: Part) -> set[str]:
    keys = {normalize(part.name)}
    keys.update(normalize(item) for item in parse_aliases(part.aliases))
    keys.update(part.aliases_norm.split(","))
    return {item for item in keys if item}


def _package_ok(part: Part, short: str) -> bool:
    if not short:
        return True
    haystack = f"{part.name_norm} {part.aliases_norm} {normalize(part.name)} {normalize(part.aliases)}"
    return short in haystack.replace(",", " ")


def _loose_text(text: str) -> str:
    return (
        text.replace("Ω", "")
        .replace("ω", "")
        .replace("µ", "u")
        .replace("μ", "u")
        .lower()
    )


def _pieces(text: str) -> list[str]:
    return [item for item in re.split(r"[\s,，()（）/]+", _loose_text(text)) if item]


def _res_values(text: str) -> set[str]:
    blob = _loose_text(text)
    found: set[str] = set()
    for match in re.finditer(r"(?<![0-9.])(\d+)([kmg])(\d+)r?(?![a-z0-9.])", blob):
        found.add(f"{match.group(1)}.{match.group(3)}{match.group(2)}")
    for match in re.finditer(r"(?<![0-9.])(\d+\.?\d*)([kmg])r?(?![a-z0-9.])", blob):
        found.add(match.group(1) + match.group(2))
    for match in re.finditer(r"(?<![0-9.])(\d+\.?\d*)r(?![a-z0-9.])", blob):
        found.add(match.group(1))
    return found


def _cap_values(text: str) -> set[str]:
    blob = _loose_text(text)
    return {
        match.group(1) + match.group(2) + "f"
        for match in re.finditer(r"(?<![0-9.])(\d+\.?\d*)([pnu])f(?![a-z0-9.])", blob)
    }


def _volt_values(text: str) -> set[str]:
    return {
        match.group(1)
        for piece in _pieces(text)
        if (match := re.fullmatch(r"(\d+\.?\d*)v", piece))
    }


def _package_codes(text: str) -> set[str]:
    return set(_PKG_RE.findall(_loose_text(text)))


def _is_precision(text: str) -> bool:
    return "0.1%" in text or "高精度" in text


def _packages_conflict(part: Part, short: str) -> bool:
    if not re.fullmatch(r"\d{4}", short or ""):
        return False
    pkgs = _package_codes(f"{part.name} {part.aliases}")
    return bool(pkgs) and short not in pkgs


def _mpn_tokens(text: str) -> set[str]:
    found: set[str] = set()
    for piece in _pieces(text):
        if len(piece) < 6 or piece in _PKG_CODES:
            continue
        if re.search(r"[\u4e00-\u9fff]", piece):
            continue
        if not re.search(r"[a-z]", piece) or not re.search(r"\d", piece):
            continue
        if re.fullmatch(r"[0-9.]+[kmg]?r?", piece) or re.fullmatch(r"[0-9.]+[pnu]f", piece):
            continue
        if re.fullmatch(r"[0-9.]+v", piece):
            continue
        found.add(piece)
    return found


def _passive_hits(line: ParsedLine, parts: list[Part], short: str) -> list[Part]:
    bom_res = _res_values(line.name)
    bom_caps = _cap_values(line.name)
    if bool(bom_res) == bool(bom_caps):
        return []
    kind = "cap" if bom_caps else "res"
    bom_volts = _volt_values(line.name)
    bom_precise = _is_precision(line.name)
    pkg = short if re.fullmatch(r"\d{4}", short or "") else ""
    strong: list[Part] = []
    weak: list[Part] = []
    for part in parts:
        if not part.active:
            continue
        text = f"{part.name} {part.aliases}"
        if _is_precision(text) and not bom_precise:
            continue
        if kind == "res":
            if not (_res_values(text) & bom_res):
                continue
            if re.search(r"ntc|热敏", text, re.I) and not re.search(r"ntc|热敏", line.name, re.I):
                continue
        else:
            if not (_cap_values(text) & bom_caps):
                continue
            volts = _volt_values(text)
            if bom_volts and volts and not (bom_volts & volts):
                continue
        pkgs = _package_codes(text)
        if pkg and pkgs and pkg not in pkgs:
            continue
        if pkg and pkg in pkgs:
            strong.append(part)
        else:
            weak.append(part)
    chosen = strong or weak
    if kind == "cap" and not bom_volts and len(chosen) > 1:
        signatures = {
            frozenset(_volt_values(f"{part.name} {part.aliases}")) for part in chosen
        }
        if len(signatures) > 1:
            return []
    if len(chosen) == 1:
        return chosen
    return []


def _mpn_hits(line: ParsedLine, parts: list[Part]) -> list[Part]:
    bom_tokens = _mpn_tokens(line.name)
    bom_whole = re.sub(r"\s+", "", _loose_text(line.name))
    hits: list[Part] = []
    for part in parts:
        if not part.active:
            continue
        matched = False
        for token in _mpn_tokens(f"{part.name} {part.aliases}"):
            if token in bom_tokens or token == bom_whole:
                matched = True
                break
            if bom_whole.startswith(token) and re.fullmatch(r"[a-z]{1,4}", bom_whole[len(token) :]):
                matched = True
                break
            targets = bom_tokens | {bom_whole}
            if any(
                token.startswith(bom) and re.fullmatch(r"-[a-z0-9]+", token[len(bom) :])
                for bom in targets
            ):
                matched = True
                break
        if matched:
            hits.append(part)
    if len(hits) == 1:
        return hits
    return []


def match_part(line: ParsedLine, parts: list[Part]) -> Part | None:
    if line.skip_bin:
        return None
    variants = name_variants(line.name)
    short = footprint_short(line.footprint)
    code = lcsc_code(line.supplier)
    named: list[Part] = []
    for part in parts:
        if not part.active:
            continue
        if _packages_conflict(part, short):
            continue
        keys = _part_keys(part)
        if keys & variants:
            named.append(part)
            if _package_ok(part, short):
                return part
    if code:
        for part in parts:
            if not part.active:
                continue
            if code in _part_keys(part):
                return part
    if len(named) == 1:
        return named[0]
    passive = _passive_hits(line, parts, short)
    if len(passive) == 1:
        return passive[0]
    mpn = _mpn_hits(line, parts)
    if len(mpn) == 1:
        return mpn[0]
    return None
